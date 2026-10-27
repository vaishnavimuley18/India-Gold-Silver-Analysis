"""STEP 0 - Build the raw CSV files from the REAL source documents (nothing is estimated).

Run:  python -m src.build_raw_from_sources

* WGC workbooks (data/source/*.xlsx) are PARSED programmatically, so every number is traceable.
* IBM publications (PDF) are TRANSCRIBED below, each value with its table reference. Re-check against the PDF if in doubt.
* Anything not in the supplied documents (silver price/demand, RBI gold stock series) is LEFT BLANK on purpose.
"""
import re
import openpyxl
import pandas as pd
from . import config as C

SRC = C.ROOT / "data" / "source"
Q2_FILE = SRC / "GDT_Tables_Q2_26_EN.xlsx"
Q1_FILE = SRC / "GDT_Tables_Q1_26_EN.xlsx"


# ---------------------------------------------------------------- WGC parsing helpers
def _quarter_cols(ws, hdr_row=5):
    return {c.column: str(c.value) for c in ws[hdr_row] if c.value and re.fullmatch(r"Q\d'\d\d", str(c.value))}


def _year_cols(ws, hdr_row=5):
    return {c.column: int(c.value) for c in ws[hdr_row] if isinstance(c.value, (int, float)) and 2000 < c.value < 2100}


def _row(ws, label, cols, hdr_row=5):
    for r in ws.iter_rows(min_row=hdr_row + 1):
        if r[1].value is not None and str(r[1].value).strip() == label:
            return {cols[c.column]: c.value for c in r if c.column in cols}
    raise KeyError(label)


def _yq(q):
    m = re.fullmatch(r"Q(\d)'(\d\d)", q)
    return 2000 + int(m.group(2)), int(m.group(1))


def wgc_quarterly(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    j, b, g = wb["Jewellery"], wb["Bar and Coin"], wb["Gold Prices"]
    jew = _row(j, "India", _quarter_cols(j))
    bar = _row(b, "India", _quarter_cols(b))
    usd = _row(g, "US$/oz", _quarter_cols(g))
    inr = _row(g, "Rs/10g", _quarter_cols(g))
    rows = []
    for q in jew:
        y, qq = _yq(q)
        rows.append(dict(Year=y, Quarter=f"Q{qq}", Gold_Price_USD_oz=usd.get(q), Gold_Price_INR_10g=inr.get(q),
                         Jewellery_Demand_t=jew.get(q), Bar_Coin_Demand_t=bar.get(q),
                         ETF_Demand_t=None, Total_Demand_t=None))
    return pd.DataFrame(rows)


def wgc_india_supply_annual(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["India Supply"]
    yc = _year_cols(ws)
    out = {k: _row(ws, k, yc) for k in ["Gross Bullion Imports", "Net Bullion Imports", "Scrap",
                                         "Domestic Supply Other Sources", "Total Supply"]}
    return pd.DataFrame(out).rename(columns={
        "Gross Bullion Imports": "Gold_Imports_t", "Net Bullion Imports": "Gold_Net_Imports_t",
        "Scrap": "Gold_Recycling_t", "Domestic Supply Other Sources": "Gold_Domestic_Other_t",
        "Total Supply": "Gold_WGC_Total_Supply_t"}).rename_axis("Year").reset_index()


# ---------------------------------------------------------------- IBM transcriptions (PDF tables)
# Gold & silver METAL production, kg. Fiscal year (April-March) labelled by its START year (2022-23 -> 2022).
# 2013-14..2020-21: IBM 'Indian Mineral Industry at a Glance 2022-23', Table 17 (gold incl. by-product recovery
#                   from imported copper; silver by-product).
# 2021-22..2023-24 : Indian Minerals Yearbook 2024 (IMYB), Table 6 'Production and Value of Non-Ferrous Metals'
#                   (2023-24 provisional). NOTE: IMIG Table 17 prints silver 2021-22 = 647,156 kg, which is the IMYB
#                   *mineral* series; the revised metal figure 713,653 kg (IMYB T6) is used so the series is consistent.
GOLD_METAL_KG = {2013: 9209, 2014: 9988, 2015: 10412, 2016: 10082, 2017: 12500, 2018: 12623, 2019: 8382,
                 2020: 7387, 2021: 10087, 2022: 13854, 2023: 16639}
SILVER_METAL_KG = {2013: 349774, 2014: 402467, 2015: 426443, 2016: 460811, 2017: 557691, 2018: 679386,
                   2019: 609340, 2020: 705796, 2021: 713653, 2022: 794431, 2023: 819725}
# Silver foreign trade, tonnes, FY start year. IMYB 2024 Table 4 (exports) and Table 11 (imports).
SILVER_IMPORTS_T = {2021: 4422, 2022: 8156, 2023: 7337}
SILVER_EXPORTS_T = {2021: 89, 2022: 130, 2023: 132}
# Gold exports, kg (IMYB 2024 Table 4 / Table 10) - fiscal basis, kept for reference only.
GOLD_EXPORTS_KG = {2021: 126, 2022: 2623, 2023: 1521}

# NMI 'At a Glance 2025' (IBM), Tables 5.15 (gold) and 5.37 (silver), as on 01.04.2025, METAL content, tonnes.
# columns: State, Mineral, Resources (A+B) , Reserves (A)
RESOURCES_2025 = [
    # ---- gold (primary metal; Kerala placer 5.42 t is listed separately, not added) ----
    ("Andhra Pradesh", "Gold", 49.66, 7.77), ("Bihar", "Gold", 37.60, 0.0), ("Chhattisgarh", "Gold", 5.81, 0.0),
    ("Jharkhand", "Gold", 16.68, 0.35), ("Karnataka", "Gold", 275.09, 66.53), ("Kerala", "Gold", 0.64, 0.0),
    ("Madhya Pradesh", "Gold", 19.98, 7.25), ("Maharashtra", "Gold", 3.64, 0.0), ("Odisha", "Gold", 1.06, 0.0),
    ("Rajasthan", "Gold", 244.91, 0.0), ("Tamil Nadu", "Gold", 1.00, 0.0), ("Uttar Pradesh", "Gold", 2.08, 0.0),
    ("West Bengal", "Gold", 0.65, 0.0),
    # ---- silver ----
    ("Andhra Pradesh", "Silver", 128.13, 0.0), ("Chhattisgarh", "Silver", 3.161, 0.0), ("Jharkhand", "Silver", 5.22, 0.0),
    ("Karnataka", "Silver", 9.72, 5.43), ("Madhya Pradesh", "Silver", 148.86, 21.86), ("Maharashtra", "Silver", 0.23, 0.0),
    ("Meghalaya", "Silver", 19.80, 0.0), ("Odisha", "Silver", 67.15, 0.0), ("Rajasthan", "Silver", 26199.65, 9541.57),
    ("Sikkim", "Silver", 56.69, 0.0), ("Tamil Nadu", "Silver", 42.55, 0.0), ("Uttarakhand", "Silver", 138.59, 0.0),
]
RESOURCE_TOTALS_2025 = {"Gold": (658.80 - 0.0, 81.90), "Silver": (26819.75, 9568.86)}  # published all-India totals (for the check)


def build():
    # ---- 1. gold quarterly ----
    gq = wgc_quarterly(Q2_FILE)
    gq.to_csv(C.RAW_FILES["gold_quarterly"], index=False)
    print("gold_quarterly_raw.csv", gq.shape)

    # ---- 2. annual supply (WGC supply + IBM production/trade) ----
    sup = wgc_india_supply_annual(Q2_FILE)
    prod = pd.DataFrame({"Year": list(GOLD_METAL_KG)})
    prod["Gold_Production_t"] = prod["Year"].map(GOLD_METAL_KG) / 1000
    prod["Silver_Production_t"] = prod["Year"].map(SILVER_METAL_KG) / 1000
    prod["Silver_Imports_t"] = prod["Year"].map(SILVER_IMPORTS_T)
    prod["Silver_Exports_t"] = prod["Year"].map(SILVER_EXPORTS_T)
    ann = sup.merge(prod, on="Year", how="outer").sort_values("Year")
    ann["Gold_Exports_t"] = ann["Year"].map(GOLD_EXPORTS_KG) / 1000
    cols = ["Year", "Gold_Production_t", "Gold_Imports_t", "Gold_Exports_t", "Gold_Recycling_t", "Silver_Production_t",
            "Silver_Imports_t", "Silver_Exports_t", "Gold_Net_Imports_t", "Gold_Domestic_Other_t", "Gold_WGC_Total_Supply_t"]
    ann[cols].to_csv(C.RAW_FILES["annual_supply"], index=False)
    print("annual_supply_raw.csv", ann.shape)

    # ---- 3. resources ----
    res = pd.DataFrame(RESOURCES_2025, columns=["State", "Mineral", "Resources_Metal_t", "Reserves_Metal_t"])
    res["Year"] = 2025
    res["Production_t"] = None
    res = res[["State", "Mineral", "Year", "Resources_Metal_t", "Reserves_Metal_t", "Production_t"]]
    res.to_csv(C.RAW_FILES["resources"], index=False)
    # reconciliation against the published all-India totals
    for m, (tot_pub, rsv_pub) in RESOURCE_TOTALS_2025.items():
        d = res[res["Mineral"] == m]
        print(f"  {m}: sum of states resources = {d['Resources_Metal_t'].sum():.2f} (published {tot_pub}), "
              f"reserves = {d['Reserves_Metal_t'].sum():.2f} (published {rsv_pub})")
    return gq, ann, res


if __name__ == "__main__":
    build()
